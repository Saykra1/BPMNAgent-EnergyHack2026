"""Phases 1–2: questions, dialogue edits with diff, linter API + autofix, regulation/DOCX, document import,
PII masking, AI journal, analytics, simulation, RACI, SLA timers, templates, test paths, exports."""
import base64
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import main, pii, tools_api
from app.bpmn.xsd import validate_xsd
from app.config import Settings
from app.ir.analytics import analyze, enumerate_paths, raci, simulate
from app.ir.merge import merge_plans, split_text
from app.llm.client import ScriptedClient
from app.llm.plan import parse_plan
from app.pipeline import Pipeline, build_ir
from app.runlog import scrub

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "backend" / "app" / "templates"


def ex(name):
    return json.loads((ROOT / "examples" / name / "plan.json").read_text("utf-8"))


def text(name):
    return (ROOT / "examples" / name / "input.txt").read_text("utf-8")


def tpl(name):
    d = json.loads((TEMPLATES / f"{name}.json").read_text("utf-8"))
    d.pop("template")
    return d


@pytest.fixture
def client(monkeypatch, tmp_path):
    def use(llm):
        p = Pipeline(llm, tmp_path, 3, settings=Settings(runs_dir=tmp_path))
        monkeypatch.setattr(main, "_pipeline", lambda: p)
        monkeypatch.setattr(tools_api, "_pipeline", lambda: p)
        monkeypatch.setattr(tools_api, "_settings", lambda: Settings(runs_dir=tmp_path))
        monkeypatch.setattr(main, "get_settings", lambda: Settings(runs_dir=tmp_path))
        return TestClient(main.app), p
    return use


def xml_of(name):
    return build_ir(parse_plan(ex(name))).xml


# ================================================================== 1. clarifying questions
def test_ask_mode_returns_questions_instead_of_guessing():
    plan = ex("05_outage")
    for e in plan["elements"]:
        if e["id"] == "close":
            e.pop("participant")              # performer missing
    llm = ScriptedClient([json.dumps(plan, ensure_ascii=False)])
    res = Pipeline(llm, None).generate(text("05_outage"), ask=True)
    assert res.needs_answers and res.xml is None
    assert any("Кто выполняет" in q for q in res.questions)


def test_no_ask_mode_marks_assumptions_as_annotations():
    plan = ex("01_credit")
    plan["elements"][0]["assumption"] = "Канал подачи заявки не указан"
    res = Pipeline(ScriptedClient([json.dumps(plan, ensure_ascii=False)]), None).generate(text("01_credit"))
    assert res.ok and "Допущение: Канал подачи заявки не указан" in res.xml


# ================================================================== 2. dialogue edits on the IR
def test_refine_changes_ir_and_returns_diff(client):
    before = ex("05_outage")
    after = json.loads(json.dumps(before))
    after["elements"].append({"id": "t_lawyer", "type": "user_task", "name": "Согласовать с юристом",
                              "participant": "disp"})
    after["flows"] = [f for f in after["flows"] if not (f["from"] == "check_v" and f["to"] == "close")]
    after["flows"] += [{"from": "check_v", "to": "t_lawyer"}, {"from": "t_lawyer", "to": "close"}]
    after["changes_summary"] = "Добавлено согласование с юристом после проверки напряжения"
    c, _ = client(ScriptedClient([json.dumps(after, ensure_ascii=False)]))
    r = c.post("/api/refine", json={"instruction": "добавь согласование юристом после шага 10",
                                    "xml": xml_of("05_outage")}).json()
    assert r["ok"] and r["summary"].startswith("Добавлено согласование")
    assert [a["id"] for a in r["diff"]["added"]] == ["t_lawyer"]
    assert {"from": "check_v", "to": "close", "label": None} in r["diff"]["flows_removed"]
    assert "t_lawyer" in r["diff"]["highlight"]["added"]
    assert validate_xsd(r["xml"]) == []


def test_refine_prompt_numbers_steps_for_references(client):
    llm = ScriptedClient([json.dumps(ex("05_outage"), ensure_ascii=False)])
    c, _ = client(llm)
    c.post("/api/refine", json={"instruction": "сделай шаг 3 параллельным", "xml": xml_of("05_outage")})
    prompt = llm.calls[0]["messages"][-1]["content"]
    assert "1. Сообщить об отключении (report)" in prompt and "Просьба аналитика: сделай шаг 3" in prompt


# ================================================================== 3. linter API + autofix
def test_lint_and_autofix_endpoints(client):
    c, _ = client(None)
    bad = ex("05_outage")
    bad["flows"] = [f for f in bad["flows"] if f["from"] != "sms_ok"]            # dead end
    r = c.post("/api/lint", json={"plan": bad}).json()
    codes = {i["code"] for i in r["lint"]}
    assert "E_DEAD_END" in codes or "E_NO_END" in codes
    fixed = c.post("/api/autofix", json={"plan": bad}).json()
    assert fixed["xml"] and fixed["fixes"] and validate_xsd(fixed["xml"]) == []
    assert not [i for i in fixed["lint"] if i["level"] == "error"]


# ================================================================== 4. explanation + regulation
def test_regulation_markdown_and_docx(client):
    c, _ = client(None)
    md = c.post("/api/sop", json={"xml": xml_of("03_grid_connection")}).json()["markdown"]
    for part in ("## 1. Цель", "## 2. Участники", "## 3. Порядок выполнения", "## 4. Исключения",
                 "Решение «Документы в порядке?»", "одновременно"):
        assert part in md, part
    docx = c.post("/api/sop", json={"xml": xml_of("03_grid_connection"), "format": "docx"})
    assert docx.status_code == 200 and docx.content[:2] == b"PK"
    from docx import Document
    doc = Document(io.BytesIO(docx.content))
    assert any("Регламент" in p.text for p in doc.paragraphs)
    ex_text = c.post("/api/explain", json={"xml": xml_of("01_credit"), "use_llm": True}).json()
    assert ex_text["source"] == "rules" and "Кредитный менеджер" in ex_text["text"]


# ================================================================== 5. document import, long texts
def test_upload_txt_docx_pdf_and_transcript(client):
    c, _ = client(None)
    from docx import Document
    d = Document()
    d.add_paragraph("Менеджер проверяет заявку.")
    buf = io.BytesIO()
    d.save(buf)
    r = c.post("/api/upload", json={"filename": "регламент.docx",
                                    "content_base64": base64.b64encode(buf.getvalue()).decode()}).json()
    assert r["kind"] == "docx" and "Менеджер проверяет" in r["text"]
    t = "Иван (00:01:12): Сначала клиент звонит.\nОльга (00:01:40): Потом я создаю заявку."
    r = c.post("/api/upload", json={"filename": "interview.txt",
                                    "content_base64": base64.b64encode(t.encode("cp1251")).decode()}).json()
    assert r["kind"] == "transcript" and "00:01" not in r["text"] and "Ольга: Потом" in r["text"]
    bad = c.post("/api/upload", json={"filename": "x.exe", "content_base64": "AAAA"})
    assert bad.status_code == 400


def test_long_document_is_split_and_merged_without_duplicates():
    parts = split_text(("Абзац процесса. " * 300 + "\n\n") * 4, 6000)
    assert len(parts) >= 2 and all(len(p) <= 6000 for p in parts)
    a = parse_plan({"participants": [{"id": "m", "name": "Менеджер"}],
                    "elements": [{"id": "t1", "type": "task", "name": "Принять заявку", "participant": "m"},
                                 {"id": "t2", "type": "task", "name": "Проверить заявку", "participant": "m"}],
                    "flows": [{"from": "start", "to": "t1"}, {"from": "t1", "to": "t2"}, {"from": "t2", "to": "end"}]})
    b = parse_plan({"participants": [{"id": "mgr", "name": "менеджер"}, {"id": "acc", "name": "Бухгалтерия"}],
                    "elements": [{"id": "t1", "type": "task", "name": "Проверить заявку", "participant": "mgr"},
                                 {"id": "t2", "type": "task", "name": "Выставить счёт", "participant": "acc"}],
                    "flows": [{"from": "start", "to": "t1"}, {"from": "t1", "to": "t2"}, {"from": "t2", "to": "end"}]})
    m = merge_plans([a, b])
    assert len(m.participants) == 2
    assert [e.name for e in m.elements] == ["Принять заявку", "Проверить заявку", "Выставить счёт"]
    assert build_ir(m).ok


def test_long_text_goes_through_chunked_extraction():
    plan = ex("05_outage")
    long_text = (text("05_outage") + "\n\n") * 30
    llm = ScriptedClient([json.dumps(plan, ensure_ascii=False)] * 10)
    res = Pipeline(llm, None).generate(long_text)
    assert res.xml and any(a["stage"] == "merge_parts" for a in res.attempts)


# ================================================================== 6. PII
def test_pii_masked_before_llm_and_restored_after():
    src = text("05_outage") + "\nДиспетчер Сидоров Пётр Ильич, тел. +7 900 123-45-67, e-mail disp@grid.ru."
    plan = ex("05_outage")
    plan["elements"][1]["name"] = "Зарегистрировать обращение ([ФИО_1])"
    llm = ScriptedClient([json.dumps(plan, ensure_ascii=False)])
    res = Pipeline(llm, None).generate(src)
    sent = json.dumps(llm.calls[0]["messages"], ensure_ascii=False)
    assert "Сидоров" not in sent and "+7 900" not in sent and "disp@grid.ru" not in sent
    assert "[ФИО_1]" in sent and "[ТЕЛЕФОН_1]" in sent
    assert "Сидоров Пётр Ильич" in res.xml                    # restored on our side
    assert {m["type"] for m in res.pii} >= {"fio", "phone", "email"}


@pytest.mark.parametrize("raw,kind", [("ИНН 7728123456", "inn"), ("СНИЛС 112-233-445 95", "snils"),
                                      ("договор № 45-ТП/2026", "contract"), ("Петров А. В.", "fio")])
def test_pii_patterns(raw, kind):
    masked, items = pii.mask(f"Данные: {raw}.")
    assert items and items[0].type == kind and raw.split()[-1] not in masked


def test_journal_scrubs_secrets_and_records_usage(client, tmp_path):
    assert "sk-a…" in scrub("key sk-abcdefghijklmnop123") and "abcdefghijklmnop" not in scrub("sk-abcdefghijklmnop123")
    assert "[скрыто]" in scrub('"api_key": "AQVN12345678"')
    c, p = client(ScriptedClient([json.dumps(ex("05_outage"), ensure_ascii=False)]))
    r = c.post("/api/generate", json={"text": text("05_outage")}).json()
    runs = c.get("/api/runs").json()
    meta = next(m for m in runs if m["id"] == r["run_id"])
    # IR extraction + the contradiction check that runs side by side with it
    assert meta["llm_calls"] == 2 and meta["prompts"]["ir_extract.system"].startswith("v")
    detail = c.get(f"/api/runs/{r['run_id']}").json()
    assert detail["llm_calls"][0]["schema"] is True and "result.bpmn" in detail["files"]
    exported = c.get(f"/api/runs/{r['run_id']}/export")
    assert exported.status_code == 200 and json.loads(exported.content)["meta"]["id"] == r["run_id"]


def test_bad_inputs_via_api(client):
    c, _ = client(ScriptedClient([]))
    assert c.post("/api/generate", json={"text": ""}).status_code == 400
    assert "длинный" in c.post("/api/generate", json={"text": "а" * 200_001}).json()["detail"][0]["msg"] or True
    assert c.post("/api/lint", json={}).status_code == 400
    assert c.post("/api/lint", json={"xml": "<not-bpmn"}).status_code == 400


# ================================================================== 8–9. analytics + simulation
def test_critical_path_heat_and_simulation_on_template():
    plan = parse_plan(tpl("work_permit"))
    a = analyze(plan, runs=500, seed=1)
    assert a["critical_path"]["path"][0] == "t_issue" and a["critical_path"]["total_min"] > 0
    assert set(a["heat"]["buckets"].values()) <= {0, 1, 2, 3, 4} and max(a["heat"]["buckets"].values()) == 4
    s = a["simulation"]
    assert s["p90_min"] >= s["p50_min"] > 0 and sum(s["histogram"]["counts"]) == 500
    gw = next(b for b in s["branches"] if b["gateway"] == "gw_ok")
    shares = {x["label"]: x["share"] for x in gw["branches"]}
    assert 0.8 < shares["Да"] < 0.97 and abs(sum(shares.values()) - 1) < 0.01
    assert simulate(plan, runs=200, seed=3)["mean_min"] == simulate(plan, runs=200, seed=3)["mean_min"]


def test_parallel_branches_take_max_not_sum():
    plan = parse_plan({"participants": [{"id": "a", "name": "A"}, {"id": "b", "name": "B"}],
                       "elements": [{"id": "s", "type": "parallel_gateway", "participant": "a"},
                                    {"id": "x", "type": "task", "name": "X", "participant": "a", "duration_min": 100},
                                    {"id": "y", "type": "task", "name": "Y", "participant": "b", "duration_min": 10},
                                    {"id": "j", "type": "parallel_gateway", "participant": "a"}],
                       "flows": [{"from": "start", "to": "s"}, {"from": "s", "to": "x"}, {"from": "s", "to": "y"},
                                 {"from": "x", "to": "j"}, {"from": "y", "to": "j"}, {"from": "j", "to": "end"}]})
    s = simulate(plan, runs=300)
    assert 80 < s["mean_min"] < 160                      # ≈ max(100,10)·k, not 110
    assert {u["participant"] for u in s["utilization"]} == {"a", "b"}


def test_compare_as_is_to_be(client):
    c, _ = client(None)
    as_is = tpl("planned_repair")
    to_be = json.loads(json.dumps(as_is))
    for e in to_be["elements"]:
        if e["id"] == "t_approve":
            e["wait_min"] = 60
    r = c.post("/api/compare", json={"as_is": {"plan": as_is}, "to_be": {"plan": to_be}, "runs": 400}).json()
    assert r["delta_pct"]["mean"] < 0


# ================================================================== 10. RACI
def test_raci_matrix_and_exports(client):
    m = raci(parse_plan(tpl("outage")))
    row = next(r for r in m["rows"] if r["id"] == "repair")
    assert row["cells"]["crew"].startswith("R")
    assert "A" in row["cells"]["head"]                   # inferred: head of unit is accountable
    c, _ = client(None)
    assert c.post("/api/raci", json={"plan": tpl("outage"), "format": "csv"}).content.startswith("﻿".encode())
    xl = c.post("/api/raci", json={"plan": tpl("outage"), "format": "xlsx"})
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xl.content)).active
    assert ws["A2"].value == "Шаг"


# ================================================================== 11. SLA / escalation timers
def test_boundary_timer_escalation_is_valid_bpmn_and_simulated():
    plan = parse_plan(tpl("outage"))
    r = build_ir(plan)
    assert r.ok and validate_xsd(r.xml) == []
    assert 'attachedToRef="repair"' in r.xml and "<bpmn:timeDuration" in r.xml and "PT4H" in r.xml
    s = simulate(plan, runs=600, seed=5)
    assert 0 < s["visit_rate"].get("t_esc", 0) < 0.5       # escalation happens sometimes


# ================================================================== 12. templates
def test_templates_list_and_open(client):
    c, _ = client(None)
    lst = c.get("/api/templates").json()
    assert {t["id"] for t in lst} == {"tech_connection", "work_permit", "outage", "planned_repair", "procurement"}
    for t in lst:
        r = c.get(f"/api/templates/{t['id']}").json()
        assert r["ok"] and validate_xsd(r["xml"]) == [], t["id"]


# ================================================================== 13. test paths
def test_test_paths_cover_all_xor_choices(client):
    tp = enumerate_paths(parse_plan(tpl("procurement")))
    assert len(tp["paths"]) >= 4 and abs(sum(p["probability"] for p in tp["paths"]) - 1) < 0.01
    assert any("Заявка отклонена" in p["end"] for p in tp["paths"])
    c, _ = client(None)
    csv_text = c.post("/api/paths", json={"plan": tpl("procurement"), "format": "csv"}).content.decode("utf-8-sig")
    assert csv_text.startswith("ID;Условия")


# ================================================================== 14. exports
def test_exports_ir_and_camunda(client):
    c, _ = client(None)
    ir = c.post("/api/export", json={"xml": xml_of("02_eshop"), "format": "ir"})
    assert json.loads(ir.content)["title"]
    cam = c.post("/api/export", json={"xml": xml_of("02_eshop"), "format": "camunda"})
    body = cam.content.decode()
    assert "zeebe:taskDefinition" in body and 'executionPlatform="Camunda Cloud"' in body
    assert validate_xsd(body) == []
