"""Process run: executes a process instance step by step over the IR.

- Automatic steps (service, script, business-rule, send tasks) run their block code against the
  process variables and produce a document about the work from the block's template.
- Human steps (task, user, manual, receive tasks) stop the token: the person enters the requested
  data and must either write a report (turned into a DOCX) or attach a file; then the run continues.
- Gateways: XOR takes the first branch whose check is true, otherwise the default branch; OR takes
  every true branch; AND splits and joins; a gateway without checks asks the analyst to choose.
- Subprocesses run their inner steps and continue when all inner tokens finish.
The whole state is a JSON document, so a run survives a server restart and can be audited.
"""
from __future__ import annotations

import datetime as _dt
import io
import re

from ..ir.graph import GATEWAY_TYPES, IRGraph
from ..ir.lint import _join_of
from ..llm.plan import Plan
from .script import ScriptError, ScriptFail, eval_check, render_template, run_script

AUTO_TYPES = {"service_task", "script_task", "business_rule_task", "send_task"}
HUMAN_TYPES = {"task", "user_task", "manual_task", "receive_task"}
MAX_VISITS = 50
MAX_MOVES = 5_000
TYPE_RU = {"service_task": "сервисная задача", "script_task": "скрипт", "business_rule_task": "бизнес-правило",
           "send_task": "отправка", "task": "задача", "user_task": "пользовательская задача",
           "manual_task": "ручная работа", "receive_task": "получение"}


class RunError(Exception):
    pass


def now() -> str:
    return _dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")


def parse_value(raw):
    """Value typed by a person: numbers and yes/no become numbers and booleans."""
    if not isinstance(raw, str):
        return raw
    s = raw.strip()
    low = s.lower()
    if low in ("да", "true", "yes"):
        return True
    if low in ("нет", "false", "no"):
        return False
    if re.fullmatch(r"-?\d{1,15}", s):
        return int(s)
    if re.fullmatch(r"-?\d{1,15}[.,]\d{1,10}", s):
        return float(s.replace(",", "."))
    return s


def is_automatic(el) -> bool:
    return el.type in AUTO_TYPES


def is_human(el) -> bool:
    return el.type in HUMAN_TYPES


class ProcessRun:
    """Wraps the JSON state of one run. All mutations go through methods that keep it consistent."""

    def __init__(self, state: dict, plan: Plan):
        self.s = state
        self.plan = plan
        self.g = IRGraph.build(plan)
        self.or_join = {}
        for e in plan.elements:
            if e.type == "inclusive_gateway" and len(self.g.out[e.id]) > 1:
                join, _ = _join_of(self.g, e.id)
                if join:
                    self.or_join[e.id] = join
        self.people = {p.id: p for p in plan.performers}
        self.roles = {p.id: p.name for p in plan.participants}

    # ------------------------------------------------------------------ creation
    @classmethod
    def new(cls, run_id: str, plan: Plan, variables: dict | None = None) -> "ProcessRun":
        state = {"id": run_id, "title": plan.title, "created": now(), "status": "running",
                 "variables": dict(variables or {}), "queue": [], "waiting": [], "joins": {},
                 "or_expected": {}, "scopes": {}, "visits": {}, "history": [], "documents": [],
                 "error": None, "seq": 0, "done_nodes": []}
        run = cls(state, plan)
        g = run.g
        starts = [s for s in g.starts() if g.out[s]] or g.starts()
        for s in starts:
            run._push(s, "root")
        run.log(None, "info", "Процесс запущен")
        return run

    # ------------------------------------------------------------------ helpers
    def _push(self, node: str, scope: str, after: bool = False):
        self.s["seq"] += 1
        self.s["queue"].append({"id": f"t{self.s['seq']}", "node": node, "scope": scope, "after": after})

    def name(self, node: str) -> str:
        return self.g.name(node)

    def log(self, node: str | None, kind: str, text: str, **extra):
        self.s["history"].append({"at": now(), "node": node, "name": self.name(node) if node else "",
                                  "kind": kind, "text": text, **extra})

    def actor(self, el) -> str:
        person = self.people.get(el.performer) if el.performer else None
        role = self.roles.get(el.participant or "", "")
        if person:
            return f"{person.name}{', ' + person.position if person.position else ''}" + (f" ({role})" if role else "")
        return role or "исполнитель не указан"

    def _scope_busy(self, scope: str) -> bool:
        return any(t["scope"] == scope for t in self.s["queue"]) or any(w["scope"] == scope for w in self.s["waiting"]) \
            or (self.s["error"] or {}).get("scope") == scope

    def _finish_token(self, scope: str):
        """A token ended inside `scope`; if it was the last one in a subprocess, continue after it."""
        if scope == "root" or self._scope_busy(scope):
            return
        info = self.s["scopes"].pop(scope, None)
        if info:
            self.log(info["sub"], "done", "Подпроцесс завершён")
            self._push(info["sub"], info["parent"], after=True)

    # ------------------------------------------------------------------ main loop
    def advance(self):
        moves = 0
        while self.s["queue"] and self.s["status"] == "running":
            moves += 1
            if moves > MAX_MOVES:
                self._error(None, "root", f"Слишком много переходов (>{MAX_MOVES}) — похоже на бесконечный цикл")
                break
            tok = self.s["queue"].pop(0)
            self._step(tok)
        if self.s["status"] == "running" and not self.s["queue"]:
            if self.s["waiting"]:
                self.s["status"] = "waiting"
            else:
                self.s["status"] = "done"
                stuck = [k for k, v in self.s["joins"].items() if v]
                if stuck:
                    self.log(None, "warning", "Процесс завершился, но слияния ждали другие ветки: "
                             + ", ".join(self.name(k.split("|", 1)[1]) for k in stuck))
                self.log(None, "end", "Процесс завершён")
        return self

    def _error(self, node, scope, message, token=None):
        self.s["status"] = "error"
        self.s["error"] = {"node": node, "scope": scope, "message": message, "token": token}
        self.log(node, "error", message)

    def _step(self, tok: dict):
        node, scope = tok["node"], tok["scope"]
        g = self.g
        typ = g.type(node)
        el = g.elements.get(node)
        if not tok.get("after"):
            # ---------------- joins
            ins = g.inc[node]
            if typ == "parallel_gateway" and len(ins) > 1:
                key = f"{scope}|{node}"
                arrived = self.s["joins"].setdefault(key, [])
                arrived.append(tok["id"])
                if len(arrived) < len(ins):
                    return
                self.s["joins"][key] = []
            elif typ == "inclusive_gateway" and len(ins) > 1:
                key = f"{scope}|{node}"
                expected = self.s["or_expected"].get(key, 1)
                arrived = self.s["joins"].setdefault(key, [])
                arrived.append(tok["id"])
                if len(arrived) < expected:
                    return
                self.s["joins"][key] = []
                self.s["or_expected"].pop(key, None)
            visits = self.s["visits"]
            visits[node] = visits.get(node, 0) + 1
            if visits[node] > MAX_VISITS:
                self._error(node, scope, f"Шаг выполнен больше {MAX_VISITS} раз — проверьте условия цикла", tok)
                return
            # ---------------- work
            if el is not None and typ == "subprocess":
                inner = [n for n in g.starts(node)]
                if inner:
                    self.s["seq"] += 1
                    sid = f"s{self.s['seq']}"
                    self.s["scopes"][sid] = {"sub": node, "parent": scope}
                    self.log(node, "started", "Подпроцесс начат")
                    for n in inner:
                        self._push(n, sid)
                    return
            elif el is not None and is_automatic(el):
                if not self._run_auto(el, scope, tok):
                    return
            elif el is not None and is_human(el):
                self.s["waiting"].append({**tok, "kind": "task", "since": now()})
                self.log(node, "waiting", f"Ожидает исполнителя: {self.actor(el)}")
                return
            elif typ == "end_event" or node == "end":
                self.log(node, "end", f"Достигнуто завершение «{self.name(node)}»" if node != "end" else "Конец ветки")
                if el is not None and el.event == "terminate":
                    self.s["queue"].clear()
                    self.s["waiting"].clear()
                    self.log(node, "info", "Терминирующее завершение: остальные ветки остановлены")
                    return
            elif typ in ("timer_event", "message_event", "message_throw_event"):
                self.log(node, "info", {"timer_event": "Таймер: при запуске ожидание пропущено",
                                        "message_event": "Сообщение считается полученным",
                                        "message_throw_event": "Сообщение отправлено"}[typ])
        self._route(node, scope, tok)

    def _route(self, node: str, scope: str, tok: dict):
        g = self.g
        typ = g.type(node)
        outs = g.out[node]
        if not outs or g.is_end(node):
            if node not in self.s["done_nodes"]:
                self.s["done_nodes"].append(node)
            self._finish_token(scope)
            return
        if typ in ("exclusive_gateway", "inclusive_gateway", "event_based_gateway") and len(outs) > 1:
            picked = self._decide(node, scope, tok)
            if picked is None:
                return
            chosen = [outs[i] for i in picked]
        else:
            chosen = outs
        if node not in self.s["done_nodes"]:
            self.s["done_nodes"].append(node)
        for f in chosen:
            self._push(f.target, scope)

    def _decide(self, node: str, scope: str, tok: dict) -> list[int] | None:
        g = self.g
        outs = g.out[node]
        typ = g.type(node)
        if typ == "event_based_gateway" or not any(f.check for f in outs):
            self.s["waiting"].append({**tok, "kind": "choice", "since": now(), "multi": typ == "inclusive_gateway",
                                      "options": [f.label or self.name(f.target) for f in outs]})
            self.log(node, "waiting", "Нужно выбрать ветку: у развилки нет проверок")
            return None
        true, details = [], []
        for i, f in enumerate(outs):
            if not f.check:
                continue
            try:
                ok = eval_check(f.check, self.s["variables"])
            except ScriptError as e:
                self._error(node, scope, f"Проверка ветки «{f.label or self.name(f.target)}»: {e}", tok)
                return None
            details.append(f"{f.check} → {'да' if ok else 'нет'}")
            if ok:
                true.append(i)
                if typ == "exclusive_gateway":
                    break
        if not true:
            default = [i for i, f in enumerate(outs) if f.default]
            unchecked = [i for i, f in enumerate(outs) if not f.check]
            true = default or (unchecked[:1] if len(unchecked) == 1 else [])
            if not true:
                self._error(node, scope, "Ни одна проверка не выполнилась, а ветки «иначе» нет: " + "; ".join(details), tok)
                return None
        labels = ", ".join(f"«{outs[i].label or self.name(outs[i].target)}»" for i in true)
        self.log(node, "chosen", f"Выбрано: {labels}", checks=details)
        if typ == "inclusive_gateway" and node in self.or_join:
            self.s["or_expected"][f"{scope}|{self.or_join[node]}"] = len(true)
        return true

    # ------------------------------------------------------------------ automatic step
    def _run_auto(self, el, scope: str, tok: dict) -> bool:
        before = dict(self.s["variables"])
        logs: list[str] = []
        if el.code.strip():
            try:
                self.s["variables"], logs = run_script(el.code, self.s["variables"])
            except ScriptFail as e:
                self._error(el.id, scope, f"Шаг сообщил об ошибке: {e.message}", tok)
                return False
            except ScriptError as e:
                self._error(el.id, scope, f"Ошибка в коде шага: {e}", tok)
                return False
        changed = {k: v for k, v in self.s["variables"].items() if before.get(k, object()) != v}
        self.log(el.id, "done", "Выполнено автоматически" + ("" if el.code.strip() else " (кода нет — шаг пропущен)"),
                 changed=changed, logs=logs)
        self._document(el, "auto", text=render_template(el.report, self.s["variables"]) if el.report else "",
                       values=changed, logs=logs)
        return True

    # ------------------------------------------------------------------ human step / choice
    def find_waiting(self, token_id: str) -> dict:
        for w in self.s["waiting"]:
            if w["id"] == token_id:
                return w
        raise RunError("Этот шаг уже не ожидает выполнения")

    def complete(self, token_id: str, text: str, values: dict, author: str = "", attachment: dict | None = None):
        if self.s["status"] not in ("waiting", "running"):
            raise RunError("Процесс не ожидает действий")
        w = self.find_waiting(token_id)
        if w["kind"] != "task":
            raise RunError("Это развилка: выберите ветку")
        el = self.g.elements[w["node"]]
        missing = [f for f in el.fields if str(values.get(f, "")).strip() == ""]
        if missing:
            raise RunError("Заполните: " + ", ".join(missing))
        if not (text or "").strip() and not attachment:
            raise RunError("Сформируйте документ о выполненной работе или прикрепите файл")
        parsed = {k: parse_value(v) for k, v in values.items() if isinstance(k, str) and k.strip()}
        self.s["variables"].update(parsed)
        self.s["waiting"].remove(w)
        self.log(el.id, "done", f"Выполнено: {author or self.actor(el)}", changed=parsed)
        self._document(el, "human", text=text or "", values=parsed, author=author or self.actor(el),
                       attachment=attachment)
        self.s["status"] = "running"
        self._route(el.id, w["scope"], w)
        return self.advance()

    def choose(self, token_id: str, picked: list[int]):
        w = self.find_waiting(token_id)
        if w["kind"] != "choice":
            raise RunError("Это не развилка")
        outs = self.g.out[w["node"]]
        picked = sorted({i for i in picked if isinstance(i, int) and 0 <= i < len(outs)})
        if not picked or (len(picked) > 1 and not w.get("multi")):
            raise RunError("Выберите одну ветку" if not w.get("multi") else "Выберите хотя бы одну ветку")
        self.s["waiting"].remove(w)
        node, scope = w["node"], w["scope"]
        self.log(node, "chosen", "Выбрано вручную: " + ", ".join(f"«{w['options'][i]}»" for i in picked))
        if self.g.type(node) == "inclusive_gateway" and node in self.or_join:
            self.s["or_expected"][f"{scope}|{self.or_join[node]}"] = len(picked)
        if node not in self.s["done_nodes"]:
            self.s["done_nodes"].append(node)
        for i in picked:
            self._push(outs[i].target, scope)
        self.s["status"] = "running"
        return self.advance()

    def retry(self):
        """After the analyst fixed the code or checks: run the failed step again."""
        err = self.s.get("error")
        if self.s["status"] != "error" or not err:
            raise RunError("Повторять нечего: процесс не остановлен ошибкой")
        tok = err.get("token")
        self.s["error"] = None
        self.s["status"] = "running"
        if tok:
            self.s["visits"][tok["node"]] = max(0, self.s["visits"].get(tok["node"], 1) - 1)
            self.s["queue"].insert(0, {"id": tok["id"], "node": tok["node"], "scope": tok["scope"],
                                       "after": tok.get("after", False)})
        self.log(err.get("node"), "info", "Повтор шага после исправления")
        return self.advance()

    # ------------------------------------------------------------------ documents
    def _document(self, el, kind: str, text: str, values: dict, logs: list | None = None, author: str = "",
                  attachment: dict | None = None):
        n = len(self.s["documents"]) + 1
        doc = {"n": n, "node": el.id, "step": el.name or el.id, "kind": kind, "created": now(),
               "author": author or ("Система (автоматически)" if kind == "auto" else self.actor(el)),
               "role": self.roles.get(el.participant or "", ""), "text": text, "values": values,
               "logs": logs or [], "file": f"{n:02d}_{_slug(el.name or el.id)}.docx", "attachment": None}
        if attachment:
            doc["attachment"] = {"name": attachment["name"], "stored": attachment["stored"], "size": attachment["size"]}
        self.s["documents"].append(doc)
        return doc

    def document_docx(self, doc: dict) -> bytes:
        from docx import Document
        from docx.shared import Pt
        d = Document()
        style = d.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        d.add_heading(f"Отчёт о выполнении шага «{doc['step']}»", level=1)
        el = self.g.elements.get(doc["node"])
        rows = [("Процесс", self.s["title"]), ("Запуск", self.s["id"]), ("Шаг", doc["step"]),
                ("Тип", TYPE_RU.get(el.type, el.type) if el else ""), ("Роль", doc["role"] or "—"),
                ("Выполнил", doc["author"]), ("Дата и время", doc["created"])]
        if el and el.deadline:
            rows.append(("Срок по регламенту", el.deadline))
        table = d.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text, table.cell(i, 1).text = k, str(v)
        if el and el.description:
            d.add_heading("Описание шага", level=2)
            d.add_paragraph(el.description)
        d.add_heading("Выполненная работа", level=2)
        d.add_paragraph(doc["text"] or ("Шаг выполнен автоматически." if doc["kind"] == "auto" else "См. приложенный файл."))
        if doc["values"]:
            d.add_heading("Данные" if doc["kind"] == "human" else "Изменённые данные процесса", level=2)
            t = d.add_table(rows=len(doc["values"]) + 1, cols=2)
            t.style = "Table Grid"
            t.cell(0, 0).text, t.cell(0, 1).text = "Показатель", "Значение"
            for i, (k, v) in enumerate(doc["values"].items(), 1):
                t.cell(i, 0).text, t.cell(i, 1).text = str(k), _fmt(v)
        if doc["logs"]:
            d.add_heading("Журнал выполнения", level=2)
            for line in doc["logs"]:
                d.add_paragraph(line, style="List Bullet")
        if doc.get("attachment"):
            d.add_heading("Приложение", level=2)
            d.add_paragraph(f"{doc['attachment']['name']} ({doc['attachment']['size'] // 1024 + 1} КБ)")
        if el and el.documents:
            d.add_paragraph("Документы по регламенту: " + ", ".join(el.documents))
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def journal_markdown(self) -> str:
        out = [f"# Журнал запуска «{self.s['title']}»", "", f"Запуск {self.s['id']}, начат {self.s['created']}, "
               f"статус: {STATUS_RU.get(self.s['status'], self.s['status'])}.", "", "| Время | Шаг | Событие |", "|---|---|---|"]
        for h in self.s["history"]:
            out.append(f"| {h['at']} | {(h['name'] or '—').replace('|', '/')} | {h['text'].replace('|', '/')} |")
        out += ["", "## Документы", ""]
        out += [f"- {d['file']} — «{d['step']}», {d['author']}"
                + (f"; приложение: {d['attachment']['name']}" if d.get("attachment") else "") for d in self.s["documents"]]
        return "\n".join(out) + "\n"

    def public(self) -> dict:
        """State for the UI (without internal bookkeeping)."""
        waiting = []
        for w in self.s["waiting"]:
            el = self.g.elements.get(w["node"])
            item = {"token": w["id"], "node": w["node"], "name": self.name(w["node"]), "kind": w["kind"], "since": w["since"]}
            if w["kind"] == "task" and el is not None:
                item.update(actor=self.actor(el), fields=el.fields, template=el.report, description=el.description,
                            deadline=el.deadline, documents=el.documents, type=el.type)
            else:
                item.update(options=w.get("options", []), multi=w.get("multi", False))
            waiting.append(item)
        return {"id": self.s["id"], "title": self.s["title"], "status": self.s["status"],
                "status_text": STATUS_RU.get(self.s["status"], ""), "created": self.s["created"],
                "variables": self.s["variables"], "waiting": waiting, "error": self.s["error"],
                "history": self.s["history"][-300:], "documents": self.s["documents"],
                "done_nodes": self.s["done_nodes"]}


STATUS_RU = {"running": "выполняется", "waiting": "ждёт исполнителя", "done": "завершён", "error": "остановлен ошибкой"}


def _fmt(v) -> str:
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, (list, dict)):
        import json
        return json.dumps(v, ensure_ascii=False)[:2000]
    return str(v)


def _slug(s: str) -> str:
    s = re.sub(r"[^\w\- ]+", "", s, flags=re.UNICODE).strip().replace(" ", "_")
    return s[:60] or "step"
