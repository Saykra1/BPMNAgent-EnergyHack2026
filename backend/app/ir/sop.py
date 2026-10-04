"""Process description and regulation (SOP) generated deterministically from the IR.

The same block structure is rendered to Markdown and to DOCX. An optional LLM "plain language"
explanation lives in pipeline/main; this module needs no model, so the regulation is always
reproducible and never contains facts that are not in the IR.
"""
from __future__ import annotations

import io

from ..llm.plan import Plan
from .graph import GATEWAY_TYPES, TASK_TYPES, IRGraph

TYPE_RU = {"user_task": "в информационной системе", "service_task": "автоматически (система)",
           "manual_task": "вручную", "send_task": "отправка", "receive_task": "ожидание ответа",
           "script_task": "автоматически (скрипт)", "business_rule_task": "по бизнес-правилу",
           "subprocess": "подпроцесс", "task": ""}


def _fmt_minutes(m: float | None) -> str:
    if m is None:
        return ""
    if m < 60:
        return f"{m:g} мин"
    if m < 60 * 24:
        return f"{m / 60:.1f} ч".replace(".0 ", " ")
    return f"{m / 60 / 8:.1f} раб. дн".replace(".0 ", " ")


def build_sop(plan: Plan) -> list[tuple]:
    """Blocks: ("h1"|"h2"|"p"|"li"|"table", payload)."""
    g = IRGraph.build(plan)
    names = {p.id: p.name for p in plan.participants}
    people = {p.id: p for p in plan.performers}
    order = [n for c in g.containers() for n in g.topo_order(c) if n in g.elements]
    num = {}
    k = 0
    for n in order:
        if g.type(n) not in GATEWAY_TYPES:
            k += 1
            num[n] = f"3.{k}"
    ref = lambda n: (f"п. {num[n]}" if n in num else ("завершение процесса" if g.is_end(n) else g.name(n)))  # noqa: E731
    back = g.back_edges()

    blocks: list[tuple] = [("h1", f"Регламент процесса «{plan.title}»")]
    blocks.append(("h2", "1. Цель и область применения"))
    org = f" в организации «{plan.organization}»" if plan.organization else ""
    blocks.append(("p", f"Регламент устанавливает порядок выполнения процесса «{plan.title}»{org}: участников, "
                        "последовательность действий, условия ветвлений, сроки и исключительные ситуации."))

    blocks.append(("h2", "2. Участники и ответственность"))
    rows = [["Участник", "Тип", "Выполняет шаги"]]
    for p in plan.participants:
        steps = [num[e.id] for e in plan.elements if e.participant == p.id and e.id in num]
        rows.append([p.name, "внешняя организация" if p.external else "внутренний участник", ", ".join(steps) or "—"])
    blocks.append(("table", rows))
    if plan.performers:
        prow = [["Исполнитель", "Должность", "Роль", "Контакты"]]
        prow += [[p.name, p.position or "—", names.get(p.role, "—"), p.contacts or "—"] for p in plan.performers]
        blocks.append(("table", prow))

    blocks.append(("h2", "3. Порядок выполнения"))
    first = [f.target for f in g.out["start"]] if g.out["start"] else []
    if first:
        blocks.append(("p", "Процесс начинается с " + ", ".join(ref(n) for n in first) + "."))
    for n in order:
        e = g.elements[n]
        if e.type in GATEWAY_TYPES:
            outs = g.out[n]
            if e.type == "parallel_gateway" and len(outs) > 1:
                blocks.append(("li", "Далее одновременно (параллельно) выполняются: "
                               + "; ".join(ref(f.target) for f in outs) + "."))
            elif e.type in ("exclusive_gateway", "inclusive_gateway") and len(outs) > 1:
                kind = "Выбирается одна ветка" if e.type == "exclusive_gateway" else "Выполняются все подходящие ветки"
                variants = []
                for f in outs:
                    cond = f.label or ("иначе" if f.default else "без условия")
                    if f.default and f.label:
                        cond += " (по умолчанию)"
                    loop = " (возврат)" if (f.source, f.target) in back else ""
                    variants.append(f"если «{cond}» — {ref(f.target)}{loop}")
                blocks.append(("li", f"Решение «{e.name or 'условие'}». {kind}: " + "; ".join(variants) + "."))
            continue
        who = names.get(e.participant, "исполнитель не указан")
        person = people.get(e.performer)
        if person:
            who += f" ({person.name}{', ' + person.position if person.position else ''})"
        line = f"{num[n]}. {e.name or e.id} — {who}"
        how = TYPE_RU.get(e.type, "")
        if how:
            line += f" ({how})"
        if e.type == "end_event":
            line = f"{num[n]}. Завершение: {e.name or 'процесс завершён'}"
        elif e.type in ("timer_event", "message_event"):
            line = f"{num[n]}. Ожидание: {e.name}"
        details = []
        if e.description:
            details.append(e.description.replace("\n", " "))
        if e.deadline:
            details.append(f"срок: {e.deadline}")
        if e.sla_hours:
            details.append(f"норматив: {e.sla_hours:g} ч")
        if e.duration_min:
            details.append(f"длительность: {_fmt_minutes(e.duration_min)}" + (" (оценка)" if e.estimate else ""))
        if e.documents:
            details.append("документы: " + ", ".join(e.documents))
        nxt = [f for f in g.out[n] if g.type(f.target) not in GATEWAY_TYPES or True]
        if len(nxt) == 1 and e.type != "end_event":
            t = nxt[0].target
            if g.type(t) not in GATEWAY_TYPES:
                loop = " (повторно)" if (n, t) in back else ""
                details.append(f"далее: {ref(t)}{loop}")
        if details:
            line += ". " + "; ".join(details)
        blocks.append(("li", line + "."))

    ends = [e for e in plan.elements if e.type == "end_event" and g.inc[e.id]]
    blocks.append(("h2", "4. Исключения и альтернативные завершения"))
    if ends:
        for e in ends:
            src = ", ".join(ref(f.source) if f.source in num else g.name(f.source) for f in g.inc[e.id])
            blocks.append(("li", f"«{e.name or 'Завершение'}» — после: {src}."))
    else:
        blocks.append(("p", "Альтернативных завершений нет: процесс заканчивается после последнего шага."))
    loops = [(s, t) for s, t in back]
    if loops:
        blocks.append(("p", "Возвраты на доработку: " + "; ".join(f"{g.name(s)} → {ref(t)}" for s, t in loops) + "."))

    timed = [e for e in plan.elements if e.deadline or e.sla_hours or e.duration_min]
    blocks.append(("h2", "5. Сроки"))
    if timed:
        rows = [["Шаг", "Срок / норматив", "Длительность"]]
        for e in timed:
            sla = "; ".join(x for x in (e.deadline, f"{e.sla_hours:g} ч" if e.sla_hours else "") if x)
            rows.append([f"{num.get(e.id, '')} {e.name}".strip(), sla or "—",
                         (_fmt_minutes(e.duration_min) + (" (оценка)" if e.estimate else "")) if e.duration_min else "—"])
        blocks.append(("table", rows))
    else:
        blocks.append(("p", "Сроки в описании процесса не указаны."))

    docs = sorted({d for e in plan.elements for d in e.documents})
    blocks.append(("h2", "6. Документы"))
    blocks.append(("p", ", ".join(docs) + "." if docs else "Документы в описании не названы."))

    blocks.append(("h2", "7. Допущения и открытые вопросы"))
    items = list(plan.assumptions) + [f"{e.name}: {e.assumption}" for e in plan.elements if e.assumption]
    for a in items:
        blocks.append(("li", "Допущение: " + a))
    for q in plan.questions:
        blocks.append(("li", "Вопрос: " + q))
    if not items and not plan.questions:
        blocks.append(("p", "Нет."))
    return blocks


def to_markdown(blocks: list[tuple]) -> str:
    out = []
    for kind, payload in blocks:
        if kind == "h1":
            out.append(f"# {payload}\n")
        elif kind == "h2":
            out.append(f"\n## {payload}\n")
        elif kind == "p":
            out.append(payload + "\n")
        elif kind == "li":
            out.append(f"- {payload}")
        elif kind == "table":
            head, *rows = payload
            out.append("| " + " | ".join(head) + " |")
            out.append("|" + "---|" * len(head))
            out += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows]
            out.append("")
    return "\n".join(out).strip() + "\n"


def to_docx(blocks: list[tuple]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    for kind, payload in blocks:
        if kind == "h1":
            doc.add_heading(payload, level=0)
        elif kind == "h2":
            doc.add_heading(payload, level=1)
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "li":
            doc.add_paragraph(payload, style="List Bullet")
        elif kind == "table":
            table = doc.add_table(rows=len(payload), cols=len(payload[0]))
            table.style = "Table Grid"
            for i, row in enumerate(payload):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = str(cell)
                    if i == 0:
                        for r in table.cell(i, j).paragraphs[0].runs:
                            r.bold = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plain_explanation(plan: Plan) -> str:
    """Model-free short explanation (used when no LLM is configured)."""
    g = IRGraph.build(plan)
    roles = [p.name for p in plan.participants if not p.external]
    ext = [p.name for p in plan.participants if p.external]
    tasks = [g.elements[n] for c in g.containers() for n in g.topo_order(c)
             if n in g.elements and g.elements[n].type in TASK_TYPES]
    xor = [e for e in plan.elements if e.type in ("exclusive_gateway", "inclusive_gateway") and len(g.out[e.id]) > 1]
    par = [e for e in plan.elements if e.type == "parallel_gateway" and len(g.out[e.id]) > 1]
    ends = [e.name for e in plan.elements if e.type == "end_event" and e.name]
    text = f"Процесс «{plan.title}»"
    text += f" выполняют: {', '.join(roles)}" if roles else ""
    text += f"; во взаимодействии с: {', '.join(ext)}" if ext else ""
    text += f". Он состоит из {len(tasks)} действий"
    if tasks:
        text += f": начинается с шага «{tasks[0].name}» и заканчивается шагом «{tasks[-1].name}»"
    text += ". "
    if xor:
        text += "Ключевые решения: " + "; ".join(f"«{e.name}»" for e in xor) + ". "
    if par:
        text += f"В процессе {len(par)} участок(а) параллельной работы. "
    if g.back_edges():
        text += "Есть возвраты на доработку. "
    if ends:
        text += "Возможные исходы: " + ", ".join(f"«{x}»" for x in ends) + "."
    return text.strip()
