"""Process analytics without real logs: durations, critical path, heat map, Monte Carlo, test paths, RACI.

Durations come from the IR (set by the analyst, or estimated by the LLM with estimate=true). Missing
values fall back to typical values per task type and are reported as "типовая оценка", never as facts.
"""
from __future__ import annotations

import csv
import io
import random
import statistics

from ..llm.plan import Plan
from .engine import Engine, NeedChoice
from .graph import GATEWAY_TYPES, TASK_TYPES, IRGraph

# typical minutes (work, wait) by task type — used only when nothing is given; flagged as defaults
TYPICAL = {"user_task": (30, 60), "task": (30, 60), "manual_task": (120, 120), "service_task": (1, 0),
           "script_task": (1, 0), "send_task": (5, 0), "receive_task": (0, 1440),
           "business_rule_task": (5, 0), "subprocess": (240, 240), "timer_event": (0, 1440),
           "message_event": (0, 480)}


def durations(plan: Plan) -> dict[str, dict]:
    """id -> {work, wait, source: given|estimate|typical}"""
    out = {}
    for e in plan.elements:
        if e.type not in TASK_TYPES and e.type not in ("timer_event", "message_event"):
            continue
        typ_work, typ_wait = TYPICAL.get(e.type, (0, 0))
        if e.duration_min is not None or e.wait_min is not None:
            out[e.id] = {"work": e.duration_min if e.duration_min is not None else typ_work,
                         "wait": e.wait_min if e.wait_min is not None else 0.0,
                         "source": "estimate" if e.estimate else "given"}
        else:
            out[e.id] = {"work": typ_work, "wait": typ_wait, "source": "typical"}
    return out


def _probabilities(g: IRGraph, gw: str) -> list[float]:
    outs = g.out[gw]
    given = [f.probability for f in outs]
    rest = 1.0 - sum(p for p in given if p is not None)
    free = [i for i, p in enumerate(given) if p is None]
    probs = [p if p is not None else max(rest, 0.0) / len(free) if free else 0.0 for p in given]
    total = sum(probs) or 1.0
    return [p / total for p in probs]


# ----------------------------------------------------------------------------- critical path
def critical_path(plan: Plan, dur: dict | None = None) -> dict:
    """Longest (worst-case) path from start to an end over the DAG (loop edges ignored)."""
    dur = dur or durations(plan)
    g = IRGraph.build(plan)
    back = g.back_edges()
    order = g.topo_order()
    best: dict[str, tuple[float, str | None]] = {}
    for s in g.starts():
        best[s] = (0.0, None)
    for n in order:
        if n not in best:
            continue
        cost = best[n][0] + (dur[n]["work"] + dur[n]["wait"] if n in dur else 0)
        for f in g.out[n]:
            if (f.source, f.target) in back:
                continue
            if f.target not in best or best[f.target][0] < cost:
                best[f.target] = (cost, n)
    ends = [n for n in best if g.is_end(n) or not g.out[n]]
    if not ends:
        return {"path": [], "total_min": 0}
    end = max(ends, key=lambda n: best[n][0] + (dur.get(n, {}).get("work", 0) + dur.get(n, {}).get("wait", 0)))
    path, cur = [], end
    while cur is not None:
        path.append(cur)
        cur = best[cur][1]
    path.reverse()
    total = sum(dur[n]["work"] + dur[n]["wait"] for n in path if n in dur)
    return {"path": [n for n in path if n in g.elements], "total_min": round(total, 1)}


def heat(plan: Plan, metric: str = "total", dur: dict | None = None) -> dict:
    """Per element value and bucket 0..4 (cold → hot) for colouring the diagram."""
    dur = dur or durations(plan)
    vals = {}
    for n, d in dur.items():
        vals[n] = {"total": d["work"] + d["wait"], "work": d["work"], "wait": d["wait"]}[metric]
    if not vals:
        return {"metric": metric, "values": {}, "buckets": {}}
    hi = max(vals.values()) or 1
    buckets = {n: min(4, int(v / hi * 5)) if hi else 0 for n, v in vals.items()}
    return {"metric": metric, "values": vals, "buckets": buckets, "max": hi}


# ----------------------------------------------------------------------------- Monte Carlo
def simulate(plan: Plan, runs: int = 1000, seed: int = 7, dur: dict | None = None) -> dict:
    dur = dur or durations(plan)
    eng = Engine(plan, max_visits=4)
    g = eng.g
    rng = random.Random(seed)
    probs = {e.id: _probabilities(g, e.id) for e in plan.elements
             if e.type in ("exclusive_gateway", "event_based_gateway", "inclusive_gateway") and len(g.out[e.id]) > 1}

    def choose(gw, options, visit):
        outs = g.out[gw]
        p = probs[gw]
        if g.type(gw) == "inclusive_gateway":
            picked = tuple(i for i in range(min(len(outs), 4)) if rng.random() < (p[i] if p[i] < 1 else 0.999))
            picked = picked or (max(range(len(p)), key=lambda i: p[i]),)
            return picked if picked in options else options[0]
        allowed = [o[0] for o in options]
        weights = [p[i] for i in allowed]
        if sum(weights) <= 0:
            return (allowed[0],)
        return (rng.choices(allowed, weights)[0],)

    def duration(node):
        d = dur.get(node)
        if not d:
            return 0.0, 0.0
        k = rng.triangular(0.8, 1.6, 1.0)          # right-skewed: delays are more likely than speed-ups
        return d["wait"] * rng.triangular(0.5, 2.0, 1.0), d["work"] * k

    totals, busy, flow_count, visits, truncated = [], {}, {}, {}, 0
    gw_passes: dict[str, int] = {}
    for _ in range(runs):
        tr = eng.run(choose, duration)
        totals.append(tr.total)
        truncated += tr.truncated
        for p, w in tr.busy.items():
            busy.setdefault(p, []).append(w / tr.total if tr.total else 0)
        for f in tr.flows:
            flow_count[f] = flow_count.get(f, 0) + 1
        for gw, _ in tr.decisions:
            gw_passes[gw] = gw_passes.get(gw, 0) + 1
        for n in set(tr.visits):
            visits[n] = visits.get(n, 0) + 1
    totals.sort()
    names = {p.id: p.name for p in plan.participants}
    branches = []
    for gw in probs:
        outs = g.out[gw]
        reached = gw_passes.get(gw, 0) or 1
        branches.append({"gateway": gw, "name": g.name(gw), "branches": [
            {"to": f.target, "to_name": g.name(f.target), "label": f.label or ("иначе" if f.default else ""),
             "share": round(flow_count.get((gw, f.target), 0) / reached, 3),
             "probability": round(probs[gw][i], 3)} for i, f in enumerate(outs)]})
    q = lambda p: totals[min(len(totals) - 1, int(p * len(totals)))]  # noqa: E731
    hist_lo, hist_hi = totals[0], totals[-1]
    width = (hist_hi - hist_lo) / 12 or 1
    hist = [0] * 12
    for t in totals:
        hist[min(11, int((t - hist_lo) / width))] += 1
    return {
        "runs": runs, "seed": seed,
        "mean_min": round(statistics.fmean(totals), 1), "p50_min": round(q(0.5), 1), "p90_min": round(q(0.9), 1),
        "min_min": round(totals[0], 1), "max_min": round(totals[-1], 1),
        "histogram": {"from": round(hist_lo, 1), "step": round(width, 1), "counts": hist},
        "utilization": sorted([{"participant": p, "name": names.get(p, p),
                                "share_of_time": round(statistics.fmean(v), 3)} for p, v in busy.items()],
                              key=lambda x: -x["share_of_time"]),
        "branches": branches,
        "visit_rate": {n: round(c / runs, 3) for n, c in visits.items() if n in g.elements},
        "truncated_runs": truncated,
        "durations_source": {n: d["source"] for n, d in dur.items()},
    }


def analyze(plan: Plan, runs: int = 1000, seed: int = 7, metric: str = "total") -> dict:
    dur = durations(plan)
    missing = [n for n, d in dur.items() if d["source"] == "typical"]
    return {"durations": dur, "critical_path": critical_path(plan, dur), "heat": heat(plan, metric, dur),
            "simulation": simulate(plan, runs, seed, dur), "typical_defaults": missing,
            "note": ("Длительности части шагов не заданы — использованы типовые значения (оценка). "
                     "Укажите длительности в карточке шага или попросите модель их оценить.") if missing else ""}


def compare(as_is: Plan, to_be: Plan, runs: int = 1000, seed: int = 7) -> dict:
    a, b = analyze(as_is, runs, seed), analyze(to_be, runs, seed)
    def pct(x, y):
        return round((y - x) / x * 100, 1) if x else None
    return {"as_is": {k: a["simulation"][k] for k in ("mean_min", "p90_min")} | {"critical_min": a["critical_path"]["total_min"],
                                                                                    "steps": sum(1 for e in as_is.elements if e.type in TASK_TYPES)},
            "to_be": {k: b["simulation"][k] for k in ("mean_min", "p90_min")} | {"critical_min": b["critical_path"]["total_min"],
                                                                                   "steps": sum(1 for e in to_be.elements if e.type in TASK_TYPES)},
            "delta_pct": {"mean": pct(a["simulation"]["mean_min"], b["simulation"]["mean_min"]),
                          "p90": pct(a["simulation"]["p90_min"], b["simulation"]["p90_min"]),
                          "critical": pct(a["critical_path"]["total_min"], b["critical_path"]["total_min"])}}


# ----------------------------------------------------------------------------- test paths
def enumerate_paths(plan: Plan, limit: int = 40) -> dict:
    """All scenarios through XOR/OR choices (loops at most once), with conditions and probability."""
    eng = Engine(plan, max_visits=1)
    g = eng.g
    dur = durations(plan)

    def duration(node):
        d = dur.get(node)
        return (d["wait"], d["work"]) if d else (0.0, 0.0)

    results, queue, truncated = [], [()], False
    while queue:
        if len(results) >= limit:
            truncated = True
            break
        prefix = queue.pop(0)
        idx = [0]

        def choose(gw, options, visit, prefix=prefix, idx=idx):
            i = idx[0]
            idx[0] += 1
            if i < len(prefix):
                return prefix[i] if prefix[i] in options else options[0]
            raise NeedChoice(gw, options)
        try:
            tr = eng.run(choose, duration)
        except NeedChoice as nc:
            for opt in nc.options:
                queue.append(prefix + (opt,))
            continue
        probs = 1.0
        conditions = []
        for gw, pick in tr.decisions:
            p = _probabilities(g, gw)
            outs = g.out[gw]
            for i in pick:
                lab = outs[i].label or ("иначе" if outs[i].default else "→ " + g.name(outs[i].target))
                conditions.append(f"{g.name(gw)}: {lab}")
            probs *= sum(p[i] for i in pick) if len(pick) == 1 else _or_prob(p, pick)
        steps = [n for n in tr.visits if n in g.elements and g.type(n) not in GATEWAY_TYPES]
        results.append({"id": f"TC{len(results) + 1:02d}", "conditions": conditions,
                        "steps": [g.name(n) for n in steps], "step_ids": steps,
                        "end": ", ".join(g.name(n) for n in tr.ends) or "—",
                        "probability": round(probs, 4), "time_min": round(tr.total, 1),
                        "loops": tr.truncated})
    results.sort(key=lambda r: -r["probability"])
    for i, r in enumerate(results):
        r["id"] = f"TC{i + 1:02d}"
    return {"paths": results, "truncated": truncated, "limit": limit}


def _or_prob(p: list[float], pick: tuple) -> float:
    out = 1.0
    for i in range(len(p)):
        out *= p[i] if i in pick else (1 - p[i])
    return out


def paths_csv(paths: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["ID", "Условия", "Шаги", "Завершение", "Вероятность", "Время, мин"])
    for r in paths["paths"]:
        w.writerow([r["id"], " | ".join(r["conditions"]), " → ".join(r["steps"]), r["end"],
                    r["probability"], r["time_min"]])
    return buf.getvalue()


def paths_markdown(paths: dict, title: str) -> str:
    lines = [f"# Тестовые сценарии: {title}", ""]
    for r in paths["paths"]:
        lines.append(f"## {r['id']} · вероятность {r['probability']:.1%} · ≈{r['time_min']:.0f} мин")
        if r["conditions"]:
            lines.append("**Условия:** " + "; ".join(r["conditions"]))
        lines.append("")
        lines += [f"{i}. {s}" for i, s in enumerate(r["steps"], 1)]
        lines.append(f"\n**Ожидаемый результат:** {r['end']}\n")
    if paths["truncated"]:
        lines.append(f"_Показаны первые {paths['limit']} сценариев._")
    return "\n".join(lines)


# ----------------------------------------------------------------------------- RACI
HEAD_WORDS = ("руковод", "начальник", "директор", "главный", "заведующ", "head", "lead")


def raci(plan: Plan) -> dict:
    """Matrix step × participant. R = performer; A/C/I from the IR, otherwise inferred and flagged."""
    g = IRGraph.build(plan)
    parts = plan.participants
    heads = [p.id for p in parts if any(w in p.name.lower() for w in HEAD_WORDS) and not p.external]
    rows = []
    for c in g.containers():
        for n in g.topo_order(c):
            e = g.elements.get(n)
            if not e or e.type not in TASK_TYPES:
                continue
            cells = {p.id: "" for p in parts}
            inferred = []
            def put(pid, letter):
                if pid in cells and letter not in cells[pid]:
                    cells[pid] = (cells[pid] + "," + letter).strip(",")
            if e.participant:
                put(e.participant, "R")
            if e.accountable:
                put(e.accountable, "A")
            elif heads and e.participant not in heads:
                put(heads[0], "A")
                inferred.append("A — руководитель (выведено)")
            elif e.participant:
                put(e.participant, "A")
            for pid in e.consulted:
                put(pid, "C")
            informed = list(e.informed)
            for f in g.out[n]:                       # hand-off: next performer is informed
                nxt = g.elements.get(f.target)
                if nxt and nxt.participant and nxt.participant != e.participant and nxt.participant not in informed:
                    informed.append(nxt.participant)
                    inferred.append("I — следующий исполнитель (выведено)")
            for mf in plan.message_flows:
                if mf.source == n and mf.target in cells:
                    informed.append(mf.target)
            for pid in informed:
                put(pid, "I")
            person = next((p for p in plan.performers if p.id == e.performer), None)
            rows.append({"id": n, "step": e.name, "cells": cells, "inferred": sorted(set(inferred)),
                         "performer": person.name if person else ""})
    return {"participants": [{"id": p.id, "name": p.name} for p in parts], "rows": rows,
            "legend": {"R": "исполняет", "A": "отвечает за результат", "C": "консультирует", "I": "информируется"}}


def raci_csv(m: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Шаг", "Исполнитель"] + [p["name"] for p in m["participants"]])
    for r in m["rows"]:
        w.writerow([r["step"], r.get("performer", "")] + [r["cells"][p["id"]] for p in m["participants"]])
    return buf.getvalue()


def raci_xlsx(m: dict, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "RACI"
    ws.append([f"RACI: {title}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Шаг"] + [p["name"] for p in m["participants"]])
    for c in ws[2]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    fills = {"R": "C6EFCE", "A": "FFEB9C", "C": "BDD7EE", "I": "EDEDED"}
    for r in m["rows"]:
        ws.append([r["step"]] + [r["cells"][p["id"]] for p in m["participants"]])
        for c in ws[ws.max_row][1:]:
            if c.value:
                c.fill = PatternFill("solid", fgColor=fills.get(str(c.value)[0], "FFFFFF"))
                c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 42
    for i in range(len(m["participants"])):
        ws.column_dimensions[chr(66 + i) if i < 25 else "Z"].width = 18
    ws.append([])
    ws.append(["R — исполняет, A — отвечает, C — консультирует, I — информируется"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
