"""Text extraction from uploaded documents: txt / md / docx / pdf (and interview transcripts)."""
from __future__ import annotations

import base64
import io
import re

MAX_BYTES = 15 * 1024 * 1024


class DocumentError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "koi8-r"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _transcript(text: str) -> str:
    """Interview transcript «Имя (00:01:02): реплика» → keep speaker and text, drop timestamps."""
    lines = []
    for line in text.splitlines():
        line = re.sub(r"^\s*\[?\d{1,2}:\d{2}(?::\d{2})?\]?\s*", "", line)
        line = re.sub(r"\s*\(\d{1,2}:\d{2}(?::\d{2})?\)\s*:", ":", line)
        lines.append(line)
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise DocumentError(f"Файл больше {MAX_BYTES // 1024 // 1024} МБ")
    if not data:
        raise DocumentError("Файл пустой")
    name = (filename or "").lower()
    if name.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
        except Exception as e:  # noqa: BLE001 - corrupted file
            raise DocumentError(f"Не удалось прочитать DOCX: {e}") from e
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text, kind = "\n".join(parts), "docx"
    elif name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:  # noqa: BLE001
            raise DocumentError(f"Не удалось прочитать PDF: {e}") from e
        kind = "pdf"
        if len(text.strip()) < 20:
            raise DocumentError("В PDF нет текстового слоя (скан). Распознайте текст (OCR) и загрузите снова.")
    elif name.endswith((".txt", ".md", ".csv", ".srt", ".vtt", "")) or "." not in name:
        text, kind = _decode_text(data), "text"
        if name.endswith((".srt", ".vtt")) or re.search(r"\(\d{1,2}:\d{2}(?::\d{2})?\)\s*:", text):
            text = re.sub(r"^\d+\s*$|^\d{2}:\d{2}:\d{2}[,.]\d{3}\s*-->.*$|^WEBVTT.*$", "", text, flags=re.M)
            text, kind = _transcript(text), "transcript"
    else:
        raise DocumentError("Поддерживаются файлы .txt, .md, .docx, .pdf и расшифровки .srt/.vtt")
    text = _clean(text)
    if not text:
        raise DocumentError("В документе не найден текст")
    return {"text": text, "kind": kind, "chars": len(text)}


def extract_base64(filename: str, content_base64: str) -> dict:
    try:
        data = base64.b64decode(content_base64, validate=False)
    except (ValueError, TypeError) as e:
        raise DocumentError("Файл повреждён при передаче") from e
    return extract_text(filename, data)
